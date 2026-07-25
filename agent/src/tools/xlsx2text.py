#!/usr/bin/env python3
# ReadDocument 的 xlsx 转换器(2026-07 文件上传调研落地实测选型):
# pandoc 3.10 的 xlsx reader 对 openpyxl 系文件(rels 用绝对 Target)直接解析失败,
# 且对正常文件也丢表头、整数渲成 1.0 —— reader 太新不可用。openpyxl 兼容面最广、值保真。
# 输出:每个工作表一节(=== Sheet: 名 ===),TSV 行;行尾空单元格剔除。
# 安全外壳与 pandoc 同款:由 registry.ts runConverter 以固定 argv 调用(无 shell、
# 白名单 env、stdout 8MB OOM 闸、60s 超时),路径已经 confinePath 钳制。
import sys

from openpyxl import load_workbook


def main() -> int:
    if len(sys.argv) != 2:
        print("usage: xlsx2text.py <file.xlsx>", file=sys.stderr)
        return 2
    try:
        wb = load_workbook(sys.argv[1], read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001 - 报错文本回喂模型作诊断
        print(f"failed to open xlsx: {e}", file=sys.stderr)
        return 1
    out = sys.stdout
    for ws in wb.worksheets:
        out.write(f"=== Sheet: {ws.title} ===\n")
        for row in ws.iter_rows(values_only=True):
            cells = ["" if v is None else str(v) for v in row]
            while cells and cells[-1] == "":
                cells.pop()
            out.write("\t".join(cells) + "\n")
        out.write("\n")
    return 0


if __name__ == "__main__":
    sys.exit(main())
