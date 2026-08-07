"""`report` —— 从 bundle 渲染交付物:三线表 docx + Methods 段 + Results 段。

## 三线表为什么是 docx 而不是图片

调研里最硬的一条痛点:经管之家 2010-12-07 的帖子「SPSS 输出的表格复制到 word 上变成图片了?」
到 2026 年**仍是该关键词的首位搜索结果** —— 十六年没被解决。
一个纯格式问题养出一整条教程内容链 + 可下载的 SPSS 样式文件 + B 站视频 + 两家 AI 创业公司的首要卖点。
所以产出必须是**Word 原生可编辑表**,不是图片、不是 Markdown 表
(A4 侧要 .docx 中文不乱码,CS 侧要 LaTeX booktabs —— 只出 Markdown 两边都不可用)。

## 「换体系」为什么是重跑 report 而不是重算

这是方案 R 三轴正交的直接兑现:方法沿方法卡长、学科沿 policy 长、**格式沿 systems 长**。
换目标体系 = 吃同一份 bundle + 新的 `--system` 重跑本 verb,**零 LLM 调用、零重算**,秒级。
返修文案里「改格式不限次」这一条,成本上就落在这里。
"""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from ..paths import ANALYSIS_DIR, BUNDLE_SUFFIX, OUT_DIR
from ..registry import load_registry
from ..report import get_system, icmje_disclosure, methods_paragraph, results_sentence
from ..report.format import fmt_ci, fmt_mean_sd, fmt_median_iqr, fmt_p, fmt_stat


class BundleSourceError(ValueError):
    """bundle 不是引擎产出的那一份 —— 拒绝渲染。

    🔴 2026-08-06 实测事故(端到端复现两次):本函数此前对 `bundle_path` **零校验**,
    executor 侧的 `StatReport` 工具也把模型给的路径原样拼进 argv。于是:

        模型用 Write 在会话根落一份手写的 `cooked.json`(照抄 bundle 结构,数字随便填)
        → `report --bundle cooked.json --system cn-core`
        → exit 0,产出真三线表 docx +「t(67.31)=4.106,p<0.001」的 Results 段
          + ICMJE 披露 +「固定种子以保证可复现」。**全程没有 analysis/ 目录,
          一次 StatRun 都没跑。** 真实 p 是 0.4653。

    根因是一句话:**D1 数字锁锁的是 bundle 的「生产者」(analysis/ 是保留区、模型写不进去),
    却没有人约束「消费者」只能吃那个位置的 bundle** —— 攻击者根本不碰保留区。

    ⚠️ 不能改用「重算 spec_hash 校验」来堵:那是**无密钥** sha256(run.py 里就是
    `hashlib.sha256(payload)`),模型自己算得出来;而且本模块只把它当文件名词干用,
    连它存不存在都不要求。要么路径闸,要么给 bundle 加进程内 HMAC —— 这里选路径闸,
    因为保留区的物理边界已经在那儿了,复用它比新引入一份密钥便宜且不会过期。
    """


def _three_line_table(bundle: dict[str, Any], sys, method, docx_path: Path) -> None:
    """python-docx 编程产**三线表**:上下粗线 + 表头细线,无竖线、无内部横线。

    三线表不是 Word 内置样式,必须逐边设置边框 —— 这正是用户手工最费时的部分。
    """
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    from docx.shared import Pt

    doc = Document()
    # 中文不乱码:显式设东亚字体(python-docx 默认只设 latin,中文会回退成宋体或方块)
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    groups = bundle.get("groups", [])
    desc = bundle.get("descriptives", {})
    # 参数/非参与统计量符号来自**方法卡**,不在这里按 id 分支(见 report/narrative.py 的说明)。
    nonparam = method.nonparametric
    stat_sym = method.statistic_symbol
    # ── C5:交付层不再裸下标 bundle["statistic"] ──────────────────────────────
    #
    # 🔴 Fisher 精确检验与 McNemar 精确检验**没有检验统计量**(方法卡的
    #    `statistic_symbol` 是 null,`reports_required` 里也没有 `statistic`)。
    #    裸下标会当场 KeyError,而这里已经是交付的最后一步 —— 用户只看到「报告生成失败」。
    #    更糟的另一半:若改成 `bundle.get("statistic", 0)`,表里会印出一个**编造的 χ²=0**。
    #    所以判据取**方法卡怎么声明**,而不是 bundle 里恰好有没有那个键:
    #    卡片说有 → 必须有(缺了是引擎缺陷,让 KeyError 炸出来);卡片说没有 → 整列不出现。
    _has_stat = bool(stat_sym) and "statistic" in method.reports_required

    # 🔴 表题与表头的**文字**来自体系包(sys),不再硬编码中文。
    #    2026-08-05 实测:选 APA 时表题写「表 1 两组均数±标准差比较」、表头写「x̄±s」,
    #    而单元格是 `52.62 (4.99)` —— 表头与数据自相矛盾。加第三套体系时不该再改这个文件。
    # 配对设计不是「两组比较」:干预前/干预后是同一批人的两个**时点**。
    # 表头写「组别」、表题写「两组」在配对表里是错的(2026-08-05 实测,配对表照写「两组」);
    # 而且配对表必须有**差值行** —— 审稿人看的就是差值及其区间,不是两个时点各自的均数。
    diff = bundle.get("descriptives_diff")
    paired = bool((bundle.get("n") or {}).get("pairs"))
    desc_word = sys.caption_desc_nonparam if nonparam else sys.caption_desc_param

    # ── C1/C5:表形按**方法卡的 input_shape** 分,不按 bundle 里恰好有什么键分 ──────
    #
    # 分类资料的表与连续量的表**列结构就不同**:描述统计不是一列(x̄±s),
    # 而是每个结局水平各一列(「有效 23(76.67)」「无效 7(23.33)」)。
    # 首版若靠 `if "table" in bundle` 来判,等于让渲染依赖上游偶然产出的键 ——
    # 而键是会变的,卡片声明不会。
    is_counts = method.input_shape == "table"
    tbl = bundle.get("table") or {}
    _paired_counts = is_counts and tbl.get("kind") == "paired"

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if is_counts:
        desc_word = sys.caption_desc_counts
        _cap = (sys.caption_counts_paired if _paired_counts
                else (sys.caption_counts_multi if len(groups) > 2 else sys.caption_counts))
        _cap = _cap.format(desc=desc_word) + " " + sys.count_pct_hint
    else:
        _cap = (sys.caption_paired if paired
                else (sys.caption_multi if len(groups) > 2 else sys.table_caption)
                ).format(desc=desc_word)
    run = title.add_run(_cap)
    run.bold = True

    # 🔴 S6:表形按**数据形状**决定,不再写死 4 列。
    #    多组表要多一列字母标记 —— 那一列在农学/食品/体育论文里**就是结论**
    #    (调研称它是「GraphPad Prism 留的结构性空白」,全场竞品都不给)。
    cld = bundle.get("cld") or {}
    # 分类表的「描述统计」占 k 列(结局水平数),连续表占 1 列。
    _lv: list[str] = [str(c) for c in (tbl.get("cols") or [])] if is_counts else []
    _n_desc = len(_lv) if is_counts else 1
    # 行标签列 + n 列 + 描述列(1 或 k) + [字母标记] + **统计量/P 共用的最后一列**。
    #
    # ⚠️ 统计量与 P **共用最后一列**是本表的既有设计:统计量写在第一个数据行,
    #    P 写在末行(末行的行首标着「P 值」)。首版重构时把它们拆成两列,
    #    连续表因此从 4 列变成 5 列 —— 多出来的那一列表头空白、数据也空,
    #    而三线表的框线是按列画的,等于凭空多一段线。这里恢复共用。
    # ⚠️ 统计量**可能不存在**(Fisher / McNemar 精确法):此时这一列只有末行的 P,
    #    表头留空(P 的名字由末行行首的「P 值」承担,与既有设计一致)。
    ncols = 2 + _n_desc + (1 if cld else 0) + 1
    _c_cld = 2 + _n_desc
    _c_stat = _c_p = ncols - 1
    # 分类表的行是列联表的行(独立样本 = 分组;配对 = 前测水平),不一定等于 `groups`。
    _rows_lab: list[str] = [str(r) for r in (tbl.get("rows") or [])] if is_counts else list(groups)
    n_extra = 1 if (paired and diff and not is_counts) else 0
    table = doc.add_table(rows=2 + len(_rows_lab) + n_extra, cols=ncols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    if is_counts:
        _row_hdr = sys.label_paired_row if _paired_counts else sys.label_group
        hdr = [_row_hdr, sys.label_n, *_lv]
    else:
        hdr = [sys.label_timepoint if paired else sys.label_group, sys.label_n,
               sys.label_desc_nonparam if nonparam else sys.label_desc_param]
    for i, h in enumerate(hdr):
        table.cell(0, i).text = h
    if cld:
        table.cell(0, _c_cld).text = sys.label_cld
    if _has_stat:
        table.cell(0, _c_stat).text = f"{stat_sym}{sys.stat_label_suffix}"
    if sys.italic_stats:
        # APA 要求统计量符号斜体(M / SD / t / p / df)。中文期刊不要求。
        # 位置随列数变(多组表多一列字母标记、精确检验没有统计量列),
        # 写死 (2,3) 会在 5 列表里给**字母标记列**加斜体而漏掉统计量列。
        # 分类表的描述列是水平名(「有效」「无效」),不是统计量符号 —— 不加斜体。
        for c in ([] if is_counts else [2]) + ([_c_stat] if _has_stat else []):
            for p_ in table.cell(0, c).paragraphs:
                for r_ in p_.runs:
                    r_.italic = True

    # 单元格逐格取数。分类表按**列联表的格子**取(row×col → n 与构成比),
    # 连续表按 descriptives 取 —— 两条路都只读上游已经算好的数,渲染层不做算术
    # (这是数字锁的前提:交付物里的每个数都必须能在 bundle 里找到源)。
    _cell_by_rc = {(str(c["row"]), str(c["col"])): c for c in (tbl.get("cells") or [])}

    for r, g in enumerate(_rows_lab, start=1):
        table.cell(r, 0).text = str(g)
        if is_counts:
            _rowcells = [_cell_by_rc.get((str(g), lv)) or {} for lv in _lv]
            table.cell(r, 1).text = str(sum(int(c.get("n", 0)) for c in _rowcells))
            for k, c in enumerate(_rowcells):
                _pct = c.get("pct")
                table.cell(r, 2 + k).text = (
                    sys.count_template.format(
                        n=int(c.get("n", 0)),
                        pct=f"{float(_pct) * 100:.{sys.count_pct_decimals}f}")
                    # `pct` 为 null = 该行一例都没有(配对表指定同一套行列时会出现)。
                    # 印 `0(nan)` 或 `0(0.00)` 都是在编数字 —— 前者是乱码,后者断言了一个
                    # 没有分母的比例。如实只印例数。
                    if _pct is not None else str(int(c.get("n", 0)))
                )
        else:
            d = desc.get(g, {})
            table.cell(r, 1).text = str(d.get("n", ""))
            table.cell(r, 2).text = (
                fmt_median_iqr(d["median"], d["q1"], d["q3"], sys) if nonparam
                else fmt_mean_sd(d["mean"], d["sd"], sys)
            )
        if cld:
            table.cell(r, _c_cld).text = str(cld.get(g, "")).strip()
        if _has_stat:
            table.cell(r, _c_stat).text = fmt_stat(bundle["statistic"], sys) if r == 1 else ""

    # 差值行:配对表的核心。参数检验给「均数差(95%CI)」,非参给中位数差与四分位 ——
    # Wilcoxon 刻意**不给均数差的 CI**:它检验的不是均数差,配套的是 Hodges-Lehmann 估计量,
    # P0 未实现。宁可不给,也不给一个与检验假设不配套的区间(会出现「区间跨 0 但 p<0.05」的自相矛盾表)。
    if paired and diff and not is_counts:
        r = len(_rows_lab) + 1
        table.cell(r, 0).text = sys.label_diff
        table.cell(r, 1).text = str(diff.get("n", ""))
        if not nonparam and "mean_ci_low" in diff:
            table.cell(r, 2).text = (
                f"{diff['mean']:.{sys.stat_decimals}f} "
                # 差值是**原始单位**的量,精度跟统计量走(不是效应量的两位)——
                # 否则同一行里点估计三位、区间两位。
                + fmt_ci(diff["mean_ci_low"], diff["mean_ci_high"], sys,
                         level=diff.get("ci_level", 0.95), decimals=sys.stat_decimals)
            )
        else:
            table.cell(r, 2).text = fmt_median_iqr(
                diff["median"], diff["q1"], diff["q3"], sys)
        table.cell(r, ncols - 1).text = ""

    last = len(_rows_lab) + 1 + n_extra
    # 🔴 原先是 `fmt_p(...).replace("p", "P")` —— **无条件**大写化。
    #    中文核心期刊惯例确实是大写 P,但 APA 明确要求**小写斜体 p**,
    #    这一行把 APA 的输出直接改错。大小写属于体系口径,由 sys.p_symbol 决定。
    table.cell(last, 0).text = f"{sys.p_symbol} 值" if sys.stat_label_suffix else sys.p_symbol
    for _c in range(1, ncols - 1):
        table.cell(last, _c).text = ""
    table.cell(last, _c_p).text = fmt_p(bundle["p_exact"], sys).replace("p", sys.p_symbol)
    if sys.italic_stats:
        for c in (0, _c_p):
            for p_ in table.cell(last, c).paragraphs:
                for r_ in p_.runs:
                    r_.italic = True

    def _border(cell, edge: str, size: int) -> None:
        tcPr = cell._tc.get_or_add_tcPr()
        borders = tcPr.find(qn("w:tcBorders"))
        if borders is None:
            borders = OxmlElement("w:tcBorders")
            tcPr.append(borders)
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        borders.append(el)

    # 三线:顶粗、表头下细、底粗。其余边一律不画 —— 有竖线就不是三线表了。
    #
    # 🔴 这里**曾经**重新写死 `ncols = 4`(2026-08-06 修)。上方已按数据形状算出
    #    `ncols = 5 if cld else 4`,而这一行把它盖回去 ⇒ 多组表的**最右列(F 值/P 值)
    #    完全没有框线**:实测该 docx 的 col4 三条边全空,而 col3 正常。
    #    症状很轻(表看起来「少一条线」),但它是三线表的**底线** —— 期刊排版要求就这三条。
    #    根因是同一个函数里对同一个量算了两次。删掉重算,复用上面那个。
    for c in range(ncols):
        _border(table.cell(0, c), "top", 12)
        _border(table.cell(0, c), "bottom", 6)
        _border(table.cell(last, c), "bottom", 12)

    doc.add_paragraph()
    note = doc.add_paragraph()
    note.add_run("注:").bold = True
    note.add_run(
        "统计方法为" + method.name_zh + ";"
        + (bundle.get("decision", {}).get("reason_zh") or "")
    )
    # 🔴 有字母标记就必须有表注解释它 —— 一列 a/ab/b 不加说明,读者无从知道它是什么。
    #    文字来自体系包(APA 与中文核心的表注写法不同),不硬编码。
    # 🔴 配对列联表的**行列各是什么**必须写在表注里。
    #    表里长这样:行首「治疗前」,而列头只有「有效 / 无效」——
    #    读者无从知道那两列是**治疗后**的水平,极容易读成「治疗前的有效/无效构成」,
    #    而那样理解的话对角线(两次都有效)就变成了一个没有意义的数。
    #    (死旋钮护栏正是从这里把 `label_paired_col` 挑出来的:声明了却没人读,
    #     等于表里少一句必须有的话。)
    if _paired_counts:
        note.add_run(
            f" 行为{sys.label_paired_row}的结局水平,列为{sys.label_paired_col}的结局水平;"
            f"对角线上的格子是两次结果一致的对子,检验只用到不一致的那两格。"
        )
    if cld:
        # 去掉体系包文案自带的「注:」前缀 —— 这里已经有一个了,不去会黏成
        # 「…事后走 Games-Howell注:同列中共享…」(实测)。
        _n = sys.note_cld
        for _pfx in ("注:", "注:", "Note. "):
            if _n.startswith(_pfx):
                _n = _n[len(_pfx):]
                break
        note.add_run(" " + _n)
    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(docx_path))


def _scipy_call_from_card(engine: str) -> str:
    """把方法卡的 `engine` 声明翻成一句可执行的 scipy 调用。

    🔴 首版在这里写了一张 `{"student_t": "...", "welch_t": "...", ...}` 的字典 ——
       被 `test_no_card_ids_in_engine_code` 当场抓住(**第四次**了)。
       方法卡里本来就有 `engine: "lib:scipy.stats.ttest_ind(equal_var=True)"`,
       在代码里再抄一份等于让「加一张方法卡」变成「还得改 report.py」,
       而方案 R 的全部意义就是让加方法只改 YAML。

    只认 `lib:` 前缀:`own:`(本仓自写的实现)没法用几行 scipy 复现,
    此时返回空串 —— **不生成脚本**,好过生成一个跑不通的。
    """
    if not engine.startswith("lib:scipy.stats."):
        return ""
    rest = engine[len("lib:scipy.stats."):]
    if "(" in rest:
        fn, arg = rest.split("(", 1)
        arg = arg.rstrip(")").strip()
        return f"stats.{fn}({{args}}, {arg})" if arg else f"stats.{fn}({{args}})"
    return f"stats.{rest}({{args}})"


def _reproduction_script(bundle: dict[str, Any], method: Any) -> str:
    """生成一份**只依赖 pandas + scipy** 的独立复算脚本。

    🔴 为什么必须有它:`interpret/SKILL.md` 让模型这样回答审稿人的「数据和代码能提供吗」——
       「报告附带的 analysis.py 就是答案(固定随机种子,可原样重跑)」。
       而 2026-08-05 之前 report 只产 docx 与 md,**analysis.py 从来不存在**。
       承诺一个不存在的交付物,比不承诺更糟:用户按图索骥找不到文件,信任一次亏光。

    🔴 为什么脚本**自带断言**:脚本算出来的数必须与报告里的数一致,否则它就成了第二个数字源 ——
       而消灭第二个数字源正是这个产品存在的理由。断言让不一致当场炸在审稿人面前,
       好过两份不同的数字各自流传。刻意**不做**「生成时偷偷校验、输出里不留断言」:
       审稿人要看见的正是这一行。

    只依赖 pandas + scipy(不需要装 statkit)—— 审稿人手里不会有我们的包。
    """
    spec = bundle.get("spec") or {}
    call = _scipy_call_from_card(getattr(method, "engine", ""))
    if not spec.get("data_file") or not spec.get("dv") or not call:
        return ""  # 老 bundle 无 spec,或该方法无法用几行 scipy 复现 —— 不硬造跑不通的脚本

    dv, group, subject = spec["dv"], spec.get("group"), spec.get("subject")
    # 统计量在不在(与三线表那侧同源判据:看方法卡怎么声明,不看 bundle 里有没有键)
    _repro_has_stat = bool(method.statistic_symbol) and "statistic" in method.reports_required
    fn = spec["data_file"]
    reader = "read_excel" if fn.lower().endswith((".xlsx", ".xlsm")) else "read_csv"
    sw = bundle.get("software", {})

    # 🔴 用引擎**实际生效**的读表参数,不用默认值(2026-08-06 修)。
    #    实测三例全崩:GBK 的中文 CSV → UnicodeDecodeError;.tsv → KeyError(按逗号切);
    #    数值编码分组 → 复算 nan。而 GBK 与「SPSS 另存为 CSV」正是我们主动指引的路径,
    #    interpret/SKILL.md 还让模型拿这个脚本回答审稿人的「数据和代码能提供吗」。
    #    一份当着审稿人的面打印「复算 nan,报告 -2.97」的脚本,比不给更糟。
    rk = spec.get("read_kwargs") or {}
    _kw = "".join(f", {k}={v!r}" for k, v in rk.items() if v not in (None, ""))
    read_line = f"df = pd.{reader}({fn!r}{_kw})"

    # 组标签在 bundle 里一律 str() 化;若原列是数值型,脚本里要按数值比较,
    # 否则 `sub[g] == '1'` 匹配不到整数 1 —— 静默得到空数组,再复算成 nan。
    _num_group = str(spec.get("group_dtype", "")).startswith(("int", "float", "uint"))

    def _lit(v: str) -> str:
        if not _num_group:
            return repr(v)
        try:
            return repr(int(v)) if float(v).is_integer() else repr(float(v))
        except (TypeError, ValueError):
            return repr(v)

    # ── 分类资料:复现的对象是**列联表**,不是若干个数值向量 ────────────────────
    #
    # 🔴 行序与列序必须从 bundle 原样固定下来(`reindex`),不能让 crosstab 自己排:
    #    crosstab 按字典序排,而引擎的列序是「事件水平在前」(见 run._order_outcome_levels)。
    #    χ² 对置换不变,所以**不会报错** —— 但率差会变号,而脚本只断言 χ² 与 p,
    #    于是一份「断言通过」的脚本会与报告里的 RD 符号相反。
    if method.input_shape == "table":
        tbl = bundle.get("table") or {}
        rows, cols = list(tbl.get("rows") or []), list(tbl.get("cols") or [])
        if spec.get("paired"):
            load = [
                f"wide = df.pivot_table(index={subject!r}, columns={group!r},",
                f"                      values={dv!r}, aggfunc='first').dropna()",
                f"tab = pd.crosstab(wide.iloc[:, 0], wide.iloc[:, 1])",
            ]
        else:
            load = [
                f"sub = df[[{dv!r}, {group!r}]].dropna()",
                f"tab = pd.crosstab(sub[{group!r}], sub[{dv!r}])",
            ]
        load += [
            f"tab = tab.reindex(index={rows!r}, columns={cols!r}, fill_value=0)",
            "print(tab)",
        ]
        names = ["tab.to_numpy()"]
    elif spec.get("paired"):
        names = ["g0", "g1"]
        load = [
            f"wide = df.pivot_table(index={subject!r}, columns={group!r},",
            f"                      values={dv!r}, aggfunc='first').dropna()",
            "g0 = wide.iloc[:, 0].to_numpy(float)",
            "g1 = wide.iloc[:, 1].to_numpy(float)",
        ]
    else:
        # 🔴 按**实际组数**生成,不写死两组(2026-08-06 修)。
        #
        #    实测事故:多组族落地后,3 组数据生成的脚本仍然是
        #      a = sub.loc[sub['g'] == 'A', ...];  b = sub.loc[sub['g'] == 'B', ...]
        #      res = stats.f_oneway(a, b)          ← 只喂了 2 组
        #      assert abs(stat - 34.09) < 1e-8     ← 而报告里是 3 组的 F
        #    真跑:「AssertionError:复算 10.537,报告 34.093」——
        #    **一份主动向审稿人声明「报告的数字复现不出来」的脚本**,比不给更糟。
        #    (上一轮修 E-6 时立的规矩正是「脚本必须真能跑」,而多组族落地时
        #     测试覆盖面没跟着扩:既有的复现脚本测试只跑两组与配对。)
        gs = list(bundle["groups"])
        names = [f"g{i}" for i in range(len(gs))]
        load = [f"sub = df[[{dv!r}, {group!r}]].dropna()"]
        load += [
            f"{nm} = sub.loc[sub[{group!r}] == {_lit(g)}, {dv!r}].to_numpy(float)"
            for nm, g in zip(names, gs)
        ]

    q = '"' * 3
    return "\n".join([
        "# -*- coding: utf-8 -*-",
        q,
        f"复算脚本 —— 报告 {bundle.get('spec_hash', '')} 的独立复现。",
        "",
        "只依赖 pandas + scipy,不需要安装 statkit。与数据文件放同一目录后直接运行。",
        "末尾的断言会核对复算结果与报告中的数值是否一致,通过即打印 ✅。",
        "",
        f"生成环境:python {sw.get('python', '?')} · numpy {sw.get('numpy', '?')} · "
        f"pandas {sw.get('pandas', '?')} · scipy {sw.get('scipy', '?')} · statkit {sw.get('statkit', '?')}",
        f"方法:{method.name_zh}({bundle['method']})",
        q,
        "import pandas as pd",
        "from scipy import stats",
        "",
        read_line,
        *load,
        "",
        f"res = {call.format(args=', '.join(names))}",
        "p = float(res.pvalue)",
        # 🔴 C5:统计量**可能不存在**。Fisher 精确检验的 `res.statistic` 是**优势比**,
        #    照着 `stat = float(res.statistic)` 写下去,脚本会打印一行
        #    「统计量 = 0.152」并拿它去和报告里根本没有的 χ² 比 —— 断言必然失败,
        #    而失败信息看起来像「引擎算错了」。判据取方法卡怎么声明,与表格那侧同源。
        *(["stat = float(res.statistic)",
           'print(f"统计量 = {stat:.6f}")'] if _repro_has_stat else []),
        'print(f"p      = {p:.10g}")',
        "",
        "# 与报告核对 —— 不一致当场报错,不让两份数字各自流传",
        *([f"assert abs(stat - {bundle['statistic']!r}) < 1e-8, \\",
           f"    f\"统计量与报告不一致:复算 {{stat}},报告 {bundle['statistic']!r}\""]
          if _repro_has_stat else []),
        f"assert abs(p - {bundle['p_exact']!r}) < 1e-10, \\",
        f"    f\"p 值与报告不一致:复算 {{p}},报告 {bundle['p_exact']!r}\"",
        'print("✅ 与报告数值完全一致")',
        "",
    ])


def _assert_engine_bundle(bundle_path: Path, workdir: Path) -> None:
    """bundle 必须落在 `<workdir>/analysis/` 下、且以 `.bundle.json` 结尾。

    用 `resolve()` 而不是词法拼接:`analysis/../cooked.json` 与
    `ln -s .. analysis/x` 两种绕法都要挡住(前者 resolve 折算掉,后者 resolve 解符号链接)。
    executor 侧 `statReportTool` 有同名的闸(走 confinePath),这里是**纵深** ——
    命令行直跑、以及将来任何新调用方,都过这一关。
    """
    root = workdir.resolve()
    target = bundle_path.resolve()
    reserved = (root / ANALYSIS_DIR).resolve()
    inside = target == reserved or reserved in target.parents
    if not inside or not target.name.endswith(BUNDLE_SUFFIX):
        raise BundleSourceError(
            f"bundle 必须是引擎产出的那一份:落在 {ANALYSIS_DIR}/ 下、以 {BUNDLE_SUFFIX} 结尾。"
            f"实得 {bundle_path!s}。\n"
            f"report 只负责把 StatRun 的结果**格式化**,它不接受在引擎之外拼出来的数字 —— "
            f"否则「交付物里的每个数字都由确定性引擎算出」这句话就不成立。\n"
            f"正确做法:先调 StatRun 得到 bundle_path,再原样把它交给 StatReport。"
        )


# ── 交付物命名(2026-08-07 · 交付链路重建 A1)─────────────────────────────────
#
# 🔴 为什么要改:交付物名字原本是 `table1_{system_id}_{spec_hash[:12]}.docx`。
#    生产实测(2026-08-07):`agent_artifacts` 1035 行里 **629 行(60.8%)文件名含 ≥8 位十六进制**,
#    而其中几乎全部来自本文件。用户在产物列表里看到的是
#      table1_cn-core_0bad84c81690.docx / table1_cn-core_442fde7bc89c.docx / table1_cn-core_cbd60ea49d6f.docx
#    —— 同一个会话里三份 37KB 的三线表,**没有任何办法判断哪份是哪个分析**。
#    这与「agent 说文件输出到 out 目录」是同一类病:引擎的内部组织逻辑泄漏到了交付面。
#    更糟的是它还会二次泄漏 —— 这三个路径键直接回给模型,模型除了照抄没有别的名字可用,
#    实测 201/221 个正文泄漏轮的燃料就是这里。
#
# 🔴 为什么不能简单换成固定中文名:那个 hash 不只是难看,它在**承担区分职责** ——
#    同一会话里的多次分析靠它区分。换成 `三线表_中文核心期刊.docx` 会让三次分析互相覆盖,
#    从「名字难看」退化成「静默丢数据」。
#
# 解法:名字 = 「产物中文名 + 方法中文名 + 体系中文名」,重名时按**中文序号**排队;
# 而序号的分配结果记在引擎保留区的 `analysis/report_names.json` 里,键是 (spec_hash, system)。
# 于是:
#   · 同一 bundle 同一体系**重复渲染 → 拿到同一个名字**(幂等:覆盖同一个文件、
#     下游 harvest 按 dedup_anchor upsert 同一行,不会每次重渲染都多一行产物);
#   · 换体系重渲染 → 体系名不同,天然是另一个名字(用户能同时留着两套格式);
#   · 不同 bundle 恰好同方法同体系 → 排到「_第2次」,仍然可读、仍然不覆盖。
#
# 台账放 analysis/(引擎保留区)而不是 out/:它是引擎的簿记,不是交付物。
NAME_LEDGER = "report_names.json"

# Windows/Explorer 与 zip 里都不安全的字符;中文与空格保留(用户要的就是可读)。
# 分隔类的 `/` `\` 换成中点而不是下划线 —— 体系名本身就长这样(「中文核心期刊 / 学位论文」),
# 换成下划线会得到「中文核心期刊 _ 学位论文」,读起来像是转义失败;中点是中文里天然的并列符。
_SEPARATOR_CHARS = "/\\|"
_UNSAFE_NAME_CHARS = ':*?"<>\r\n\t'


def _safe_name_part(s: str, *, limit: int = 40) -> str:
    """把方法名/体系名收成可安全做文件名的一段。空 → 'unknown'。"""
    out = []
    for ch in (s or ""):
        if ch in _SEPARATOR_CHARS:
            out.append("·")
        elif ch in _UNSAFE_NAME_CHARS:
            out.append(" ")
        else:
            out.append(ch)
    cleaned = "".join(out)
    # 折叠「空格·空格」→「·」与连续空格,免得中点两侧留下转义痕迹。
    while "  " in cleaned:
        cleaned = cleaned.replace("  ", " ")
    cleaned = cleaned.replace(" ·", "·").replace("· ", "·")
    cleaned = cleaned.strip("._· ")
    return (cleaned[:limit] or "unknown")


def _assign_delivery_stem(
    *, workdir: Path, spec_hash: str, system_id: str, base: str
) -> str:
    """给 (spec_hash, system) 分配一个稳定的交付物主名(不含扩展名)。

    幂等:同一键第二次调用返回同一个值。不同键撞 base 时排「_第N次」。
    台账读写失败一律**降级为直接用 base**(命名是体验问题,不该因为簿记坏了而让渲染失败)。
    """
    ledger_path = workdir / ANALYSIS_DIR / NAME_LEDGER
    key = f"{spec_hash}|{system_id}"
    ledger: dict[str, str] = {}
    try:
        if ledger_path.exists():
            loaded = json.loads(ledger_path.read_text(encoding="utf-8"))
            if isinstance(loaded, dict):
                ledger = {str(k): str(v) for k, v in loaded.items()}
    except Exception:
        ledger = {}  # 台账损坏 → 当作空,下面重新分配(最坏结果是排到 _第2次,不丢数据)

    if key in ledger:
        return ledger[key]

    taken = set(ledger.values())
    stem = base
    n = 2
    while stem in taken:
        stem = f"{base}_第{n}次"
        n += 1
        if n > 50:  # 病态兜底:极不可能,但别死循环
            stem = f"{base}_{spec_hash[:8]}"
            break

    ledger[key] = stem
    try:
        ledger_path.parent.mkdir(parents=True, exist_ok=True)
        ledger_path.write_text(json.dumps(ledger, ensure_ascii=False, indent=2), encoding="utf-8")
    except Exception:
        pass  # 写不下 → 本次仍用分配到的名字;下次可能重排,可接受
    return stem


def _merge_manifest(out_dir: Path, entries: list[dict[str, str]]) -> None:
    """把交付物的友好标题写进 out/manifest.json(后端据此给产物卡片配标题)。

    合并而非覆盖:同一会话里多次分析各写各的条目,按 file 去重、后写者胜。
    失败静默 —— 标题是锦上添花,绝不能因为它让渲染失败。

    ⚠️ manifest.json 本身**不是交付物**,后端四个 feature 都靠 agentartifact.IsManifest
    把它排除在产物列表外(2026-08-07 之前只有 claw 排了)。
    """
    path = out_dir / "manifest.json"
    existing: list[dict[str, str]] = []
    try:
        if path.exists():
            loaded = json.loads(path.read_text(encoding="utf-8"))
            if isinstance(loaded, dict) and isinstance(loaded.get("artifacts"), list):
                existing = [e for e in loaded["artifacts"] if isinstance(e, dict)]
    except Exception:
        existing = []
    by_file = {str(e.get("file", "")): e for e in existing if e.get("file")}
    for e in entries:
        by_file[e["file"]] = e
    try:
        path.write_text(
            json.dumps({"artifacts": list(by_file.values())}, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
    except Exception:
        pass


def run_report(*, bundle_path: Path, system_id: str, workdir: Path) -> dict[str, Any]:
    # 🔴 第一件事就是验来源。放在 read_text 之前:连读都不该读一份来路不明的 bundle。
    _assert_engine_bundle(bundle_path, workdir)
    sys_spec = get_system(system_id)
    bundle = json.loads(bundle_path.read_text(encoding="utf-8"))
    reg = load_registry()
    method = reg.methods.get(bundle["method"])
    if method is None:
        raise ValueError(
            f"bundle 里的方法 {bundle['method']!r} 不在注册表内 —— "
            f"多半是卡片被删/改名而旧 bundle 仍在。report 拒绝凭空猜测它的呈现口径。"
        )

    out_dir = workdir / OUT_DIR
    out_dir.mkdir(parents=True, exist_ok=True)
    # 交付物主名:中文可读 + 同会话不撞 + 同 bundle 幂等。见上方 _assign_delivery_stem 的长注释。
    stem = _assign_delivery_stem(
        workdir=workdir,
        spec_hash=bundle.get("spec_hash", "result").split(":")[-1],
        system_id=system_id,
        base=f"{_safe_name_part(method.name_zh)}_{_safe_name_part(sys_spec.name_zh)}",
    )

    docx_path = out_dir / f"三线表_{stem}.docx"
    _three_line_table(bundle, sys_spec, method, docx_path)

    # 效应量卡:决定正文里写哪个名字与符号(Cohen's d_z / d_z)。找不到卡就回落到 id,
    # 但 doctor 与 test_every_effect_id_the_engine_can_emit_has_a_card 保证这种情况不出现。
    eff_id = (bundle.get("effect") or {}).get("name")
    eff_card = reg.effects.get(eff_id) if eff_id else None
    methods = methods_paragraph(bundle, sys_spec, method, eff_card)
    results = results_sentence(bundle, sys_spec, method, eff_card)
    disclosure = icmje_disclosure(bundle)

    md_path = out_dir / f"Results段_{stem}.md"
    md_path.write_text(
        f"# 统计结果({sys_spec.name_zh})\n\n"
        f"## Methods · 统计分析\n\n{methods}\n\n"
        f"## Results\n\n{results}\n\n"
        f"## 数据分析工具披露(ICMJE)\n\n{disclosure}\n",
        encoding="utf-8",
    )

    extra: dict[str, Any] = {}
    repro = _reproduction_script(bundle, method)
    if repro:
        py_path = out_dir / f"复现脚本_{stem}.py"
        py_path.write_text(repro, encoding="utf-8")
        extra["analysis_py"] = str(py_path.relative_to(workdir))

    # 🔴 给**作者**的提醒,刻意不混进 Methods 正文 —— 正文是要粘进论文的,
    #    而这是一句写给用户看的待办。引擎只能为「它读到的那个文件」背书,
    #    合并/清洗/聚合/换算若发生在 StatRun 之外,只有用户知道,必须由用户补进 Methods。
    authoring_notes = [
        "⚠️ Methods 里的缺失计数只覆盖**喂给引擎的那个数据文件**。"
        "若这个文件是经过多表合并、清洗、聚合或单位换算后生成的,那些步骤引擎看不见 —— "
        "请把它们(排除了哪些、依据是什么、原始 n 是多少)一并写进 Methods,否则审稿人核对 n 时会对不上。",
    ]

    # ── p 值与效应量置信区间**互相矛盾**时,必须说出来 ────────────────────────
    #
    # 🔴 实测(2026-08-07,12 vs 12 的并发症数据):Fisher 精确检验 p=0.069 →
    #    Results 写「差异无统计学意义」,而同一句里印着「RD=0.42,95%CI [0.05, 0.67]」
    #    —— **区间不含 0**。两句话在同一行里互相打脸,审稿人一定会问。
    #
    # 这不是引擎算错:精确检验保守、Newcombe 区间不保守,两套程序本来就可能不一致
    # (χ² 与 Wald 区间同理)。但「本来就可能」不是不说的理由 ——
    # 沉默的后果是用户把这一行原样投出去,然后在审稿意见里第一次听说这件事。
    #
    # 处置与本仓一贯的口径一致:**不改数、不选边、如实披露**,并给出该怎么写。
    _e = bundle.get("effect") or {}
    _lo, _hi = _e.get("ci_low"), _e.get("ci_high")
    _p = bundle.get("p_exact")
    if _e.get("ci_method") != "none" and _lo is not None and _hi is not None and _p is not None:
        _alpha = ((bundle.get("multiplicity") or {}).get("alpha_nominal")) or 0.05
        _ci_excludes_null = (float(_lo) > 0) or (float(_hi) < 0)
        if (float(_p) >= _alpha) and _ci_excludes_null:
            authoring_notes.append(
                f"⚠️ 本次结果里 {sys_spec.p_symbol}={float(_p):.4g}(未达 {_alpha} 水平),"
                f"而效应量的 {int(float(_e.get('ci_level', 0.95)) * 100)}% 置信区间"
                f" [{float(_lo):.4g}, {float(_hi):.4g}] **不包含 0** —— 两者结论方向不一致。"
                "这不是计算错误:检验与区间用的是两套不同的程序(精确检验偏保守,"
                "而该区间不保守),小样本时二者可以不一致。"
                "**不要只报其中一个** —— 建议在 Results 里两个都给,并在 Discussion 里"
                "写明样本量有限、结论需谨慎;审稿人看到只报一个会认为是选择性报告。"
            )
        elif (float(_p) < _alpha) and not _ci_excludes_null:
            authoring_notes.append(
                f"⚠️ 本次结果里 {sys_spec.p_symbol}={float(_p):.4g}(达显著性水平),"
                f"而效应量的置信区间 [{float(_lo):.4g}, {float(_hi):.4g}] **包含 0** —— "
                "两者结论方向不一致,理由同上(检验与区间是两套程序)。"
                "两个都要报,并说明这一不一致。"
            )

    # 产物卡片的友好标题/说明:引擎知道方法与体系,给得出比文件名更好的说法。
    # 覆盖生产上 97.7% 的产物(agent_artifacts 里 1011/1035 行是 stat_lab)。
    _n_total = sum(int(v) for v in (bundle.get("n", {}).get("by_group") or {}).values()) or None
    _desc = f"{method.name_zh}" + (f" · n={_n_total}" if _n_total else "")
    # 🔴 同一会话里多次分析必须能在卡片上被区分开 —— 这正是用户原本的抱怨
    #    (三份 table1_cn-core_<hash>.docx 分不清哪份是哪个)。如果标题全叫「三线表(中文核心期刊)」,
    #    问题只是从文件名搬到了标题上。stem 里已经带了排队编号(_第2次),把它接到标题末尾。
    _ord = ""
    if "_第" in stem and stem.endswith("次"):
        _ord = "·" + stem[stem.rindex("_第") + 1:]
    _entries = [
        {"file": docx_path.name, "title": f"三线表({sys_spec.name_zh}){_ord}", "description": _desc},
        {"file": md_path.name, "title": f"Results / Methods 段({sys_spec.name_zh}){_ord}", "description": _desc},
    ]
    if "analysis_py" in extra:
        _entries.append({
            "file": Path(extra["analysis_py"]).name,
            "title": f"可复现分析脚本{_ord}",
            "description": f"{method.name_zh} · 独立运行可复算上表全部数字",
        })
    _merge_manifest(out_dir, _entries)

    return {
        "ok": True,
        "authoring_notes_zh": authoring_notes,
        "system": system_id,
        "system_name_zh": sys_spec.name_zh,
        "table_docx": str(docx_path.relative_to(workdir)),
        "results_md": str(md_path.relative_to(workdir)),
        **extra,
        "methods_zh": methods,
        "results_zh": results,
    }


def cli(argv: list[str]) -> tuple[int, str]:
    opts: dict[str, str] = {}
    i = 0
    while i < len(argv):
        if argv[i].startswith("--") and i + 1 < len(argv):
            opts[argv[i][2:]], i = argv[i + 1], i + 2
        else:
            i += 1
    if "bundle" not in opts or "system" not in opts:
        return 2, json.dumps(
            {"ok": False, "reason": "usage",
             "hint_zh": "用法:report --bundle <path> --system <cn-core|apa> [--workdir .]"
                        " —— **不设默认体系**:p 前导零 / 离散度符号 / 星号阈值在各体系间两两冲突。"},
            ensure_ascii=False)
    try:
        out = run_report(
            bundle_path=Path(opts["bundle"]),
            system_id=opts["system"],
            workdir=Path(opts.get("workdir", ".")),
        )
        return 0, json.dumps(out, ensure_ascii=False, indent=2)
    except BundleSourceError as e:
        # 单独一档 reason:模型要能从回执里读出「换个路径重试没用,得先跑 StatRun」,
        # 而不是把它当成一次偶发的 report_failed 去重试(实测模型会连试 3 次)。
        return 2, json.dumps(
            {"ok": False, "reason": "bundle_not_from_engine", "why_zh": str(e),
             "hint_zh": "调 StatRun 拿到 bundle_path 后原样传给 StatReport;"
                        "不要自己拼 bundle,也不要把它写到别的目录。"},
            ensure_ascii=False, indent=2)
    except Exception as e:  # noqa: BLE001
        return 2, json.dumps({"ok": False, "reason": "report_failed", "error": str(e)},
                             ensure_ascii=False, indent=2)
